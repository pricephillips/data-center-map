#!/usr/bin/env python3
"""
build_report_templates.py

Builds the two Word templates of spec 010 from code, in the Hawthorn styles:

  templates/location_report.docx   the client-neutral location report that
                                   render_location_report.py fills with docxtpl
  templates/reference.docx         the Pandoc reference document md_to_docx.sh
                                   converts markdown briefs against

Why code and not a file saved in Word: a builder is reviewable in a diff and
reproducible, Word splits Jinja tags across runs, and the selftest can prove
the template carries no typed number (Principle VI: every number in a report
is read at run time). The section order comes from templates/manifest.json,
which records Price's confirmation; client templates under
templates/clients/<client>/ are separate files that bind the same fields
(specs/010-deliverable-generation/contracts/template_fields.md).

Reads
  templates/manifest.json
Writes
  templates/location_report.docx
  templates/reference.docx

Usage
  python scripts/build_report_templates.py            rebuild both
  python scripts/build_report_templates.py --check    fail if the committed
                                                      templates differ in fields
                                                      or section order
  python scripts/build_report_templates.py --selftest
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

MANIFEST = os.path.join(ROOT, "templates", "manifest.json")
LOCATION_TEMPLATE = os.path.join(ROOT, "templates", "location_report.docx")
REFERENCE_DOCX = os.path.join(ROOT, "templates", "reference.docx")

# Hawthorn styles (research R12)
FONT = "Georgia"
INK = "1F2937"
INK_MUTED = "4B5563"
ACCENT = "1F4E79"
RULE = "D1D5DB"

TAG_RE = re.compile(r"\{%.*?%\}|\{\{.*?\}\}", re.S)
DIGIT_RE = re.compile(r"\d")


def load_manifest(path: str = MANIFEST) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


# --------------------------------------------------------------------------
# styles
# --------------------------------------------------------------------------

def hawthorn_styles(doc) -> None:
    """One style set for both templates; Pandoc's style names included."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))

    styles = doc.styles

    def font(style, size, color=INK, bold=None, italic=None):
        f = style.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), RGBColor.from_string(color)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)
        for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(theme), None)
        if bold is not None:
            f.bold = bold
        if italic is not None:
            f.italic = italic

    def para(name, size, *, base="Normal", color=INK, bold=None, italic=None, before=0,
             after=6, keep=False, align=None):
        st = styles[name] if name in [s.name for s in styles] else styles.add_style(
            name, WD_STYLE_TYPE.PARAGRAPH)
        if name != "Normal":
            st.base_style = styles[base]
        font(st, size, color, bold, italic)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = keep
        if align is not None:
            pf.alignment = align
        return st

    para("Normal", 10.5, after=6)
    para("Body Text", 10.5, after=6)
    para("First Paragraph", 10.5, base="Body Text", after=6)
    para("Compact", 10.5, base="Body Text", after=2)
    para("Title", 22, color=INK, bold=False, after=4)
    para("Subtitle", 11, color=INK_MUTED, after=12)
    para("Author", 10.5, color=INK_MUTED, after=2)
    para("Date", 10.5, color=INK_MUTED, after=12)
    para("Abstract", 10.5, italic=True, after=12)
    para("Heading 1", 15, color=ACCENT, bold=False, before=18, after=6, keep=True)
    para("Heading 2", 12, color=INK, bold=True, before=12, after=4, keep=True)
    para("Heading 3", 10.5, color=INK, bold=True, before=8, after=2, keep=True)
    para("Block Text", 10.5, color=INK_MUTED, italic=True, after=6)
    para("Caption", 9.5, color=INK_MUTED, italic=True, before=4, after=2)
    para("Image Caption", 9.5, base="Caption", color=INK_MUTED, italic=True, after=2)
    para("Table Caption", 9.5, base="Caption", color=INK_MUTED, italic=True, after=4)
    para("Credit", 8.5, color=INK_MUTED, after=1)
    para("Figure", 10.5, after=0, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Captioned Figure", 10.5, base="Figure", after=0, keep=True)
    para("Table Text", 9, after=0)
    para("Footer", 8, color=INK_MUTED, after=0)
    para("Footnote Text", 9, color=INK_MUTED, after=2)
    para("Draft Notice", 10.5, color="B91C1C", bold=True, after=6)

    names = [s.name for s in styles]
    hl = styles["Hyperlink"] if "Hyperlink" in names else styles.add_style(
        "Hyperlink", WD_STYLE_TYPE.CHARACTER)
    font(hl, 9, ACCENT)
    hl.font.underline = True
    vc = styles["Verbatim Char"] if "Verbatim Char" in names else styles.add_style(
        "Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    vc.font.name, vc.font.size = "Courier New", Pt(9)
    if "Table" not in names:
        tbl = styles.add_style("Table", WD_STYLE_TYPE.TABLE)
        tbl.base_style = styles["Table Grid"]


def _fixed_core(doc, title: str, when: str) -> None:
    cp = doc.core_properties
    stamp = dt.datetime.fromisoformat(when + "T00:00:00")
    cp.title, cp.author, cp.last_modified_by = title, "Hawthorn Intelligence", "Hawthorn Intelligence"
    cp.created = cp.modified = stamp
    cp.revision = 1
    cp.comments = "Built by scripts/build_report_templates.py (spec 010)."


def _cell(cell, text: str, *, bold: bool = False, style: str = "Table Text") -> None:
    from docx.shared import RGBColor

    cell.text = ""
    p = cell.paragraphs[0]
    p.style = style
    run = p.add_run(text)
    run.bold = bold
    if bold:
        run.font.color.rgb = RGBColor.from_string(INK)


def _loop_table(doc, headers: list[str], fields: list[str], loop: str, widths: list[float]):
    """Header row, then a {%tr for %} block of one data row (docxtpl idiom)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    t = doc.add_table(rows=4, cols=len(headers))
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, bold=True)
    _cell(t.rows[1].cells[0], "{%tr for row in " + loop + " %}")
    for i, f in enumerate(fields):
        _cell(t.rows[2].cells[i], f)
    _cell(t.rows[3].cells[0], "{%tr endfor %}")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    hdr = t.rows[0]._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    hdr.append(rep)
    return t


def _p(doc, text: str = "", style: str = "Body Text"):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    return p


def _tag(doc, tag: str):
    """A paragraph holding only a block tag; docxtpl removes it on render."""
    return _p(doc, tag, "Body Text")


def _footer(doc) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = doc.sections[0].footer.paragraphs[0]
    p.style = doc.styles["Footer"]
    p.add_run("{{ report.title }}. Data as of {{ report.as_of }}. Commit {{ report.commit_short }}. Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "-"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


# --------------------------------------------------------------------------
# the location report template
# --------------------------------------------------------------------------

def _section_title(doc, m: dict) -> None:
    _p(doc, "{{ report.title }}", "Title")
    _tag(doc, "{%p if report.draft %}")
    _p(doc, "Draft for review. Not for distribution.", "Draft Notice")
    _tag(doc, "{%p endif %}")
    _p(doc, "FIPS {{ report.fips }}. Data as of {{ report.as_of }}. "
            "Commit {{ report.commit_short }}.", "Subtitle")
    # Plate slot (spec 013 contract): full text width, directly under the
    # title block; caption and credits from the sidecar, never composed here.
    _tag(doc, "{%p if plate %}")
    _p(doc, "{{ plate.image }}", "Figure")
    _p(doc, "{{ plate.title }} {{ plate.subtitle }}", "Caption")
    _tag(doc, "{%p for line in plate.credits %}")
    _p(doc, "{{ line }}", "Credit")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")


def _section_summary(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _tag(doc, "{%p if history.enacted %}")
    _p(doc, "{{ report.county_name }} has an {{ define('enacted_restriction') }} on record.")
    _tag(doc, "{%p else %}")
    _p(doc, "No {{ define('enacted_restriction') }} is on record for {{ report.county_name }}, "
            "{{ report.state_name }}.")
    _tag(doc, "{%p endif %}")
    _p(doc, "Its {{ define('calibrated_score') }} is {{ score.calibrated }}, and its "
            "{{ define('interval') }} runs from {{ score.interval_low }} to "
            "{{ score.interval_high }}. Its {{ define('decile') }} is {{ score.decile }}.")
    _p(doc, "Each {{ define('tracked_case') }} in this report carries its own source. "
            "Tracked cases in the county: {{ cases.count }}. Adjacent counties: "
            "{{ neighbors.count }}. Adjacent counties with an enacted restriction on record: "
            "{{ neighbors.n_enacted }}.")


def _section_history(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _tag(doc, "{%p if history.enacted %}")
    _tag(doc, "{%p if history.rows %}")
    _p(doc, "Enacted restrictions recorded in the tracker, each with its own source:")
    _loop_table(doc, ["Date", "Place", "Level", "Type", "Status", "Source"],
                ["{{ row.date }}", "{{ row.place }}", "{{ row.level }}", "{{ row.type }}",
                 "{{ row.status }}", "{{r row.source }}"],
                "history.rows", [0.9, 1.5, 0.6, 1.0, 0.8, 1.7])
    _tag(doc, "{%p endif %}")
    _tag(doc, "{%p if history.census_rows %}")
    _p(doc, "Instruments filed in the external restriction census:")
    _loop_table(doc, ["Date enacted", "Instrument", "Status", "Source"],
                ["{{ row.date }}", "{{ row.instrument }}", "{{ row.status }}", "{{r row.source }}"],
                "history.census_rows", [1.0, 1.3, 1.0, 3.2])
    _tag(doc, "{%p endif %}")
    _tag(doc, "{%p if history.unitemized %}")
    _p(doc, "{{ sources.unitemized_notice }}")
    _tag(doc, "{%p endif %}")
    _tag(doc, "{%p else %}")
    _p(doc, "No enacted restriction is on record for {{ report.county_name }}, "
            "{{ report.state_name }}.")
    _tag(doc, "{%p endif %}")


def _section_score(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _p(doc, "The {{ define('calibrated_score') }} for {{ report.county_name }} is "
            "{{ score.calibrated }}. Its {{ define('interval') }} runs from "
            "{{ score.interval_low }} to {{ score.interval_high }}. "
            "{{ sources.interval_notice }}")
    _p(doc, "Its {{ define('decile') }} is {{ score.decile }}. It sits at the "
            "{{ score.pct_national }} {{ define('pct_national') }} and the "
            "{{ score.pct_state }} {{ define('pct_state') }} within {{ report.state_name }}.")
    _p(doc, "Most similar counties compared: {{ score.peer_n }}. Their "
            "{{ define('peer_median') }} is {{ score.peer_median }}.")
    _tag(doc, "{%p if charts %}")
    _p(doc, "{{ charts.score_distribution }}", "Figure")
    _p(doc, "{{ charts.score_distribution_title }}", "Caption")
    _tag(doc, "{%p endif %}")


def _section_cases(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _tag(doc, "{%p if cases.rows %}")
    _p(doc, "Each case shows its {{ define('outcome_grade') }} as recorded and its source.")
    _loop_table(doc, ["Date", "Place", "Type", "Status", "Outcome", "Project", "Source"],
                ["{{ row.date }}", "{{ row.place }}", "{{ row.type }}", "{{ row.status }}",
                 "{{ row.outcome }}", "{{ row.project }}", "{{r row.source }}"],
                "cases.rows", [0.8, 1.2, 0.9, 0.7, 1.0, 0.9, 1.0])
    _tag(doc, "{%p else %}")
    _p(doc, "No tracked case is on record for {{ report.county_name }}.")
    _tag(doc, "{%p endif %}")


def _section_neighbors(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _tag(doc, "{%p if neighbors.rows %}")
    _p(doc, "Counties that share a boundary with {{ report.county_name }}:")
    _loop_table(doc, ["County", "State", "Enacted restriction on record", "Decile", "Tracked cases"],
                ["{{ row.county }}", "{{ row.state }}", "{{ row.enacted }}", "{{ row.decile }}",
                 "{{ row.cases }}"],
                "neighbors.rows", [2.0, 0.7, 1.6, 0.8, 1.4])
    _tag(doc, "{%p else %}")
    _p(doc, "No adjacent county is listed for {{ report.county_name }}.")
    _tag(doc, "{%p endif %}")


def _section_sources(doc, heading: str) -> None:
    doc.add_heading(heading, level=1)
    _p(doc, "{{ sources.notice }}")
    _p(doc, "Every number in this report was read from the files below at commit "
            "{{ report.commit_short }}, with data as of {{ report.as_of }}. Before the report is "
            "sent, a re-derivation check compares each number with the cell it came from.")
    _loop_table(doc, ["File", "File hash (first twelve characters)"],
                ["{{ row.path }}", "{{ row.sha }}"], "sources.files", [4.0, 2.5])


SECTIONS = {
    "title": _section_title, "summary": _section_summary, "history": _section_history,
    "score": _section_score, "cases": _section_cases, "neighbors": _section_neighbors,
    "sources": _section_sources,
}


def build_location(manifest: dict):
    import docx

    doc = docx.Document()
    hawthorn_styles(doc)
    _fixed_core(doc, "Location report", manifest["confirmed"]["date"])
    for s in manifest["sections"]:
        if s["id"] not in SECTIONS:
            raise ValueError(f"manifest section {s['id']!r} has no builder")
        fn = SECTIONS[s["id"]]
        fn(doc, manifest) if s["id"] == "title" else fn(doc, s["heading"])
    body = doc.element.body
    first = body[0]
    if first.tag.endswith("}p") and not "".join(first.itertext()).strip():
        body.remove(first)
    _footer(doc)
    return doc


def build_reference(manifest: dict):
    """Pandoc reads only the styles; the sample text shows each one."""
    import docx

    doc = docx.Document()
    hawthorn_styles(doc)
    _fixed_core(doc, "Hawthorn reference document", manifest["confirmed"]["date"])
    _p(doc, "Title", "Title")
    _p(doc, "Subtitle", "Subtitle")
    _p(doc, "Author", "Author")
    _p(doc, "Date", "Date")
    _p(doc, "Abstract", "Abstract")
    for level in (1, 2, 3):
        doc.add_heading(f"Heading {level}", level=level)
    _p(doc, "First Paragraph", "First Paragraph")
    _p(doc, "Body Text", "Body Text")
    _p(doc, "Compact", "Compact")
    _p(doc, "Block Text", "Block Text")
    _p(doc, "Image Caption", "Image Caption")
    _p(doc, "Table Caption", "Table Caption")
    return doc


def to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def write(path: str, data: bytes) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(data)


# --------------------------------------------------------------------------
# inspection, shared with render_location_report.py
# --------------------------------------------------------------------------

def docx_xml_text(data: bytes) -> str:
    """All visible template text (body, headers, footers) with tags kept."""
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in sorted(z.namelist()):
            if re.match(r"word/(document|header\d*|footer\d*)\.xml$", name):
                xml = z.read(name).decode("utf-8")
                for para in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
                    out.append("".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", para)))
    return "\n".join(out)


def typed_digits(text: str) -> list[str]:
    """Digits in literal template text, outside Jinja tags."""
    return [ln for ln in TAG_RE.sub("", text).splitlines() if DIGIT_RE.search(ln)]


def template_fields(path_or_bytes) -> set[str]:
    from docxtpl import DocxTemplate

    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    return set(DocxTemplate(src).get_undeclared_template_variables())


def section_headings(data: bytes) -> list[str]:
    import docx

    d = docx.Document(io.BytesIO(data))
    return [p.text for p in d.paragraphs if p.style.name.lower() == "heading 1"]


# --------------------------------------------------------------------------
# selftest
# --------------------------------------------------------------------------

def _selftest() -> int:
    checks: list[tuple[str, bool]] = []

    def check(name: str, ok) -> None:
        checks.append((name, bool(ok)))

    m = load_manifest()
    loc = to_bytes(build_location(m))
    ref = to_bytes(build_reference(m))
    text = docx_xml_text(loc)
    check("template carries no typed digit (Principle VI)", not typed_digits(text))
    fields = template_fields(loc)
    check("template binds every required field", set(m["required_fields"]) <= fields)
    check("template binds no field outside the contract",
          fields <= set(m["required_fields"]) | set(m["optional_fields"]))
    check("sections follow the confirmed order",
          section_headings(loc) == [s["heading"] for s in m["sections"] if s["heading"]])
    check("plate slot sits directly under the title block",
          text.index("{%p if plate %}") < text.index(m["sections"][1]["heading"]))
    check("history has the plain no-restriction sentence",
          "No enacted restriction is on record for" in text)
    check("RichText fields use the r form", "{{r row.source }}" in text
          and "{{ row.source }}" not in text)
    check("tables carry an explicit tblGrid", b"<w:tblGrid>" in zipfile.ZipFile(
        io.BytesIO(loc)).read("word/document.xml"))
    check("template carries no paraId at or above 0x80000000",
          not re.search(rb'w14:paraId="[89A-Fa-f]', zipfile.ZipFile(io.BytesIO(loc)).read(
              "word/document.xml")))
    import docx

    rd = docx.Document(io.BytesIO(ref))
    names = {s.name for s in rd.styles}
    need = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Body Text",
            "First Paragraph", "Compact", "Hyperlink", "Image Caption", "Table Caption",
            "Block Text", "Author", "Date", "Abstract", "Table"}
    check("reference document carries Pandoc's style names", need <= names)
    check("Hawthorn font on body and headings",
          rd.styles["Body Text"].font.name == FONT and rd.styles["Heading 1"].font.name == FONT)
    if os.path.exists(LOCATION_TEMPLATE):
        with open(LOCATION_TEMPLATE, "rb") as fh:
            committed = fh.read()
        check("committed template matches the builder (fields and sections)",
              template_fields(committed) == fields
              and section_headings(committed) == section_headings(loc)
              and docx_xml_text(committed) == text)
    bad = dict(m, sections=m["sections"] + [{"id": "nope", "heading": "X"}])
    try:
        build_location(bad)
        check("unknown manifest section refused", False)
    except ValueError:
        check("unknown manifest section refused", True)

    fails = 0
    for name, ok in checks:
        print(("PASS " if ok else "FAIL ") + name)
        fails += 0 if ok else 1
    print(f"{len(checks) - fails}/{len(checks)} checks passed")
    return 1 if fails else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    if args.selftest:
        return _selftest()
    m = load_manifest()
    loc, ref = to_bytes(build_location(m)), to_bytes(build_reference(m))
    if args.check:
        bad = []
        for path, data in ((LOCATION_TEMPLATE, loc), (REFERENCE_DOCX, ref)):
            if not os.path.exists(path):
                bad.append(f"{os.path.relpath(path, ROOT)} missing")
                continue
            with open(path, "rb") as fh:
                old = fh.read()
            if docx_xml_text(old) != docx_xml_text(data):
                bad.append(f"{os.path.relpath(path, ROOT)} differs from the builder; rebuild it")
        for b in bad:
            print(b, file=sys.stderr)
        return 1 if bad else 0
    write(LOCATION_TEMPLATE, loc)
    write(REFERENCE_DOCX, ref)
    print(f"wrote {os.path.relpath(LOCATION_TEMPLATE, ROOT)} and {os.path.relpath(REFERENCE_DOCX, ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
