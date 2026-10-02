#!/usr/bin/env python3
"""
render_location_report.py

Fills the location report Word template with one county's facts (spec 010
US1). The base template, templates/location_report.docx, is client-neutral;
--client NAME selects templates/clients/NAME/location_report.docx instead,
which binds the same fields, so data binding never changes per client. Every
number comes from report_data.py at run time (Principle VI); the template and
the definitions carry none. Price reviews the output before it is sent.

Before anything is written, the rendered document must pass:
  the template contract    required fields present, no unknown field, no
                           refused column, no typed digit
  vocabulary and em-dash   leak_audit's blocking pattern over all text
  first-use definitions    every defined term carries its definition the
                           first time it appears (FR-004)
  the no-rate rule         a county with no enacted restriction gets one plain
                           sentence in its history section and no number
  FR-002 XML rules         paraId below 0x80000000, explicit tblGrid,
                           hyperlinks beside runs
  validate.py --original   when DOCX_VALIDATE names the docx skill's validator

--plate places a terrain plate (spec 013) in the template's fixed slot under
the title block, with the sidecar's caption and credits; the contract and its
refusals are in specs/010-deliverable-generation/plan.md.

Reads
  everything report_data.py reads; templates/manifest.json; the template;
  outputs/plates/<id>.png and .json (with --plate)
Writes
  outputs/location_reports/<stem>.docx
  outputs/location_reports/<stem>.facts.json
  deliverables/location_reports/<stem>.md

Usage
  python scripts/render_location_report.py --fips 13255 [--client NAME]
         [--plate outputs/plates/13255_spalding.png] [--draft] [--no-charts]
  python scripts/render_location_report.py --verify outputs/location_reports/<stem>.facts.json
  python scripts/render_location_report.py --selftest
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from urllib.parse import urlparse

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import build_report_templates as brt  # noqa: E402
import export_geolibre as eg  # noqa: E402
import report_data as rd  # noqa: E402

MANIFEST = "templates/manifest.json"
BASE_TEMPLATE = "templates/location_report.docx"
CLIENT_DIR = "templates/clients"
OUT_DIR = "outputs/location_reports"
MD_DIR = "deliverables/location_reports"
CLIENT_RE = re.compile(r"^[a-z0-9_-]+$")
SLUG_RE = re.compile(r"[^a-z0-9]+")
IDENT_RE = re.compile(r"[A-Za-z_]\w*")

W = rd.W
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# --------------------------------------------------------------------------
# template selection and contract
# --------------------------------------------------------------------------

def load_manifest(root: str = ROOT) -> dict:
    with open(os.path.join(root, MANIFEST), encoding="utf-8") as fh:
        return json.load(fh)


def template_path(client: str, root: str = ROOT) -> str:
    if not client:
        return os.path.join(root, BASE_TEMPLATE)
    if not CLIENT_RE.match(client):
        raise rd.ReportError(f"client name must match {CLIENT_RE.pattern}: {client!r}")
    path = os.path.join(root, CLIENT_DIR, client, "location_report.docx")
    if not os.path.exists(path):
        raise rd.ReportError(f"no client template at {os.path.relpath(path, root)}")
    return path


def check_contract(data: bytes, manifest: dict, where: str) -> None:
    """The field contract every template, base or client, is held to."""
    fields = brt.template_fields(data)
    missing = sorted(set(manifest["required_fields"]) - fields)
    unknown = sorted(fields - set(manifest["required_fields"]) - set(manifest["optional_fields"]))
    if missing:
        raise rd.ReportError(f"{where}: template lacks required field(s) {missing}")
    if unknown:
        raise rd.ReportError(f"{where}: template names field(s) outside the contract {unknown}")
    text = brt.docx_xml_text(data)
    refused = sorted({w for tag in brt.TAG_RE.findall(text) for w in IDENT_RE.findall(tag)
                      if rd.is_refused(w)})
    if refused:
        raise rd.ReportError(f"{where}: template reads refused column(s) {refused}")
    digits = brt.typed_digits(text)
    if digits:
        raise rd.ReportError(f"{where}: template text carries a typed number: {digits[0]!r}")


def text_width_emu(data: bytes) -> int:
    import docx

    sec = docx.Document(io.BytesIO(data)).sections[0]
    return int(sec.page_width - sec.left_margin - sec.right_margin)


# --------------------------------------------------------------------------
# context
# --------------------------------------------------------------------------

def link(tpl, url: str, text: str | None = None):
    from docxtpl import RichText

    rt = RichText()
    url = (url or "").strip()
    if url.lower().startswith(("http://", "https://")):
        host = urlparse(url).netloc.lower()
        host = host[4:] if host.startswith("www.") else host
        rt.add(text or host, url_id=tpl.build_url_id(url), style="Hyperlink")
    else:
        rt.add(text or url or "Not recorded")
    return rt


def context(tpl, facts: dict, *, plate: dict | None, chart_png: bytes | None,
            chart_title: str, width_emu: int) -> dict:
    from docx.shared import Emu
    from docxtpl import InlineImage

    history = dict(facts["history"])
    history["rows"] = [dict(r, source=link(tpl, r["url"])) for r in history["rows"]]
    history["census_rows"] = [dict(r, source=link(tpl, r["source"])) for r in history["census_rows"]]
    cases = dict(facts["cases"])
    cases["rows"] = [dict(r, source=link(tpl, r["url"])) for r in cases["rows"]]
    ctx = {
        "report": facts["report"], "history": history, "score": facts["score"],
        "cases": cases, "neighbors": facts["neighbors"], "sources": facts["sources"],
        "define": rd.Definer(facts["definitions"]), "plate": None, "charts": None,
    }
    if plate:
        # Full text width, aspect ratio kept (only the width is set), bytes unmodified.
        ctx["plate"] = {"image": InlineImage(tpl, plate["png"], width=Emu(width_emu)),
                        "title": plate["title"], "subtitle": plate["subtitle"],
                        "credits": plate["credits"]}
    if chart_png:
        ctx["charts"] = {"score_distribution": InlineImage(tpl, io.BytesIO(chart_png),
                                                           width=Emu(width_emu)),
                         "score_distribution_title": chart_title}
    return ctx


# --------------------------------------------------------------------------
# reading the rendered document back
# --------------------------------------------------------------------------

def _rels(z: zipfile.ZipFile) -> dict:
    xml = z.read("word/_rels/document.xml.rels").decode("utf-8")
    return dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', xml)) | {
        k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', xml)}


def blocks(data: bytes) -> list[dict]:
    """Body content in order: {'kind': 'p', 'style', 'text', 'md'} or a table."""
    from lxml import etree

    z = zipfile.ZipFile(io.BytesIO(data))
    rels = _rels(z)
    root = etree.fromstring(z.read("word/document.xml"))
    styles = {}
    sx = etree.fromstring(z.read("word/styles.xml"))
    for st in sx.iter(f"{{{W}}}style"):
        name = st.find(f"{{{W}}}name")
        styles[st.get(f"{{{W}}}styleId")] = name.get(f"{{{W}}}val") if name is not None else ""

    def para(p) -> dict:
        txt, md = [], []
        for el in p:
            if el.tag == f"{{{W}}}r":
                t = "".join(x.text or "" for x in el.iter(f"{{{W}}}t"))
                txt.append(t)
                md.append(t)
            elif el.tag == f"{{{W}}}hyperlink":
                t = "".join(x.text or "" for x in el.iter(f"{{{W}}}t"))
                url = rels.get(el.get(f"{{{R_NS}}}id"), "")
                txt.append(t)
                md.append(f"[{t}]({url})" if url else t)
        sid = p.find(f"{{{W}}}pPr/{{{W}}}pStyle")
        style = styles.get(sid.get(f"{{{W}}}val"), "") if sid is not None else "Normal"
        style = style[:1].upper() + style[1:]  # Word stores built-ins as "heading 1"
        return {"kind": "p", "style": style, "text": "".join(txt), "md": "".join(md),
                "image": p.find(f".//{{{W}}}drawing") is not None}

    out = []
    for el in root.find(f"{{{W}}}body"):
        if el.tag == f"{{{W}}}p":
            out.append(para(el))
        elif el.tag == f"{{{W}}}tbl":
            rows = []
            for tr in el.iter(f"{{{W}}}tr"):
                cells = []
                for tc in tr.iter(f"{{{W}}}tc"):
                    ps = [para(p) for p in tc.iter(f"{{{W}}}p")]
                    cells.append((" ".join(p["text"] for p in ps), " ".join(p["md"] for p in ps)))
                rows.append(cells)
            out.append({"kind": "table", "rows": rows})
    return out


def plain_text(bl: list[dict]) -> str:
    lines = []
    for b in bl:
        if b["kind"] == "p":
            lines.append(b["text"])
        else:
            lines.extend(" | ".join(c[0] for c in row) for row in b["rows"])
    return "\n".join(lines)


def section_text(bl: list[dict], heading: str) -> str | None:
    """Text from a Heading 1 to the next one; None if the heading is absent."""
    out, inside = [], False
    for b in bl:
        if b["kind"] == "p" and b["style"] == "Heading 1":
            if inside:
                break
            inside = b["text"].strip() == heading
            continue
        if inside:
            out.append(plain_text([b]))
    return "\n".join(out) if inside or out else None


def to_markdown(bl: list[dict]) -> str:
    md = []
    for b in bl:
        if b["kind"] == "table":
            rows = [[c[1].replace("|", "\\|") for c in r] for r in b["rows"]]
            if not rows:
                continue
            md.append("| " + " | ".join(rows[0]) + " |")
            md.append("|" + "|".join("---" for _ in rows[0]) + "|")
            md.extend("| " + " | ".join(r) + " |" for r in rows[1:])
            md.append("")
            continue
        text = b["md"].strip()
        if not text:
            continue
        if b["style"] == "Title":
            md += [f"# {text}", ""]
        elif b["style"] == "Heading 1":
            md += [f"## {text}", ""]
        elif b["style"] == "Heading 2":
            md += [f"### {text}", ""]
        elif b["style"] in ("Caption", "Credit", "Subtitle"):
            md += [f"*{text}*", ""]
        else:
            md += [text, ""]
    return "\n".join(md).rstrip() + "\n"


def post_checks(data: bytes, facts: dict, manifest: dict, plate: dict | None = None) -> list[dict]:
    bl = blocks(data)
    text = plain_text(bl)
    # The plate caption and credits are printed verbatim from the sidecar (spec
    # 013 contract), so the first-use rule cannot ask them to carry definitions.
    own = text
    if plate:
        for line in [f"{plate['title']} {plate['subtitle']}", *plate["credits"]]:
            own = own.replace(line, "")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml_parts = [z.read(n) for n in z.namelist()
                     if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        footer_text = " ".join(re.sub(r"<[^>]+>", "", x.decode("utf-8")) for x in xml_parts[1:])
    rd.check_text(text + "\n" + footer_text, where="rendered report")
    rd.check_first_use(own, facts["definitions"])
    for x in xml_parts:
        rd.check_docx_xml(x)
    if not facts["history"]["enacted"]:
        heading = next(s["heading"] for s in manifest["sections"] if s["id"] == "history")
        sec = section_text(bl, heading)
        if sec is None:  # a client template renamed the heading: check the sentence itself
            sec = next((b["text"] for b in bl if b["kind"] == "p"
                        and "no enacted restriction is on record" in b["text"].lower()), "")
        rd.check_no_rate(sec)
    return bl


# --------------------------------------------------------------------------
# render
# --------------------------------------------------------------------------

def stem_for(facts: dict, client: str, draft: bool) -> str:
    r = facts["report"]
    slug = SLUG_RE.sub("-", r["county_name"].lower().replace(" county", "")).strip("-")
    return f"{r['fips']}_{slug}" + (f"-{client}" if client else "") + ("-draft" if draft else "")


def run_validator(docx_bytes: bytes, original: str) -> None:
    validator = os.environ.get("DOCX_VALIDATE", "").strip()
    if not validator:
        return
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.docx")
        with open(path, "wb") as fh:
            fh.write(docx_bytes)
        r = subprocess.run([sys.executable, validator, path, "--original", original],
                           capture_output=True, text=True)
    if r.returncode != 0:
        raise rd.ReportError("validate.py --original failed:\n" + (r.stdout + r.stderr).strip())


def render(fips: str, *, root: str = ROOT, client: str = "", plate: str | None = None,
           draft: bool = False, charts: bool = True, allow_uncommitted: bool = False,
           sha: str | None = None) -> dict:
    from docxtpl import DocxTemplate

    t0 = time.time()
    manifest = load_manifest(root)
    tpath = template_path(client, root)
    with open(tpath, "rb") as fh:
        tbytes = fh.read()
    check_contract(tbytes, manifest, os.path.relpath(tpath, root))
    facts = rd.build(fips, root, sha=sha, allow_uncommitted=allow_uncommitted,
                     draft=draft, client=client)
    plate_meta = rd.load_plate(plate, draft=draft, root=root) if plate else None

    chart_png, chart_title = None, ""
    if charts:
        import charts as ch

        out = ch.render(ch.build_chart("score_distribution", fips, root, facts))
        chart_png = out["png"]
        chart_title = json.loads(out["spec"])["title"]["text"]

    tpl = DocxTemplate(io.BytesIO(tbytes))
    ctx = context(tpl, facts, plate=plate_meta, chart_png=chart_png, chart_title=chart_title,
                  width_emu=text_width_emu(tbytes))
    tpl.render(ctx, autoescape=True)
    buf = io.BytesIO()
    tpl.save(buf)
    data = buf.getvalue()
    bl = post_checks(data, facts, manifest, plate_meta)
    run_validator(data, tpath)

    stem = stem_for(facts, client, draft)
    sidecar = {
        "fips": facts["report"]["fips"], "generator": "scripts/render_location_report.py",
        "commit": facts["report"]["commit"], "as_of": facts["report"]["as_of"],
        "template": os.path.relpath(tpath, root), "template_sha256": rd.sha256(tpath),
        "manifest_tag": manifest["tag"], "client": client, "draft": draft,
        "plate": ({"png": os.path.relpath(plate_meta["png"], root),
                   "sidecar": os.path.relpath(plate_meta["sidecar"], root),
                   "commit_sha": plate_meta["commit_sha"], "preview": plate_meta["preview"]}
                  if plate_meta else None),
        "inputs": [{"path": p, "sha256": rd.sha256(os.path.join(root, p))} for p in rd.INPUTS],
        "facts": facts["facts"],
    }
    docx_path = os.path.join(root, OUT_DIR, stem + ".docx")
    facts_path = os.path.join(root, OUT_DIR, stem + ".facts.json")
    md_path = os.path.join(root, MD_DIR, stem + ".md")
    write_outputs(docx_path, data, facts_path, sidecar, md_path, to_markdown(bl))
    return {"docx": docx_path, "facts": facts_path, "md": md_path,
            "seconds": round(time.time() - t0, 2), "bytes": len(data)}


def write_outputs(docx_path: str, data: bytes, facts_path: str, sidecar: dict,
                  md_path: str, md: str) -> None:
    for p in (docx_path, md_path):
        os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(docx_path, "wb") as fh:
        fh.write(data)
    with open(facts_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(sidecar, fh, indent=2)
        fh.write("\n")
    with open(md_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(md)


# --------------------------------------------------------------------------
# re-derivation (SC-002)
# --------------------------------------------------------------------------

def verify(facts_path: str, root: str = ROOT) -> tuple[list[str], list[str]]:
    """(mismatches, notes). Zero mismatches is the pass."""
    with open(facts_path, encoding="utf-8") as fh:
        side = json.load(fh)
    docx_path = facts_path[: -len(".facts.json")] + ".docx"
    with open(docx_path, "rb") as fh:
        text = plain_text(blocks(fh.read()))
    mismatches, notes = [], []
    rebuilt = None
    tables: dict[str, dict] = {}
    for f in side["facts"]:
        path = os.path.join(root, f["file"])
        if not os.path.exists(path):
            mismatches.append(f"{f['name']}: {f['file']} is missing")
            continue
        if rd.sha256(path) != f["sha256"]:
            notes.append(f"{f['file']} changed since the render (data growth, not a mismatch)")
        if f["rule"] == "count":
            if rebuilt is None:
                rebuilt = {x["name"]: x for x in rd.build(side["fips"], root,
                                                          sha=side["commit"])["facts"]}
            got = rebuilt.get(f["name"], {}).get("formatted")
        else:
            if f["file"] not in tables:
                with open(path, encoding="utf-8-sig", newline="") as fh:
                    tables[f["file"]] = {r[f["key_column"]].zfill(5): r for r in csv.DictReader(fh)}
            row = tables[f["file"]].get(f["key"])
            got = rd.fmt(row[f["column"]], f["rule"]) if row and row.get(f["column"]) else None
        if got != f["formatted"]:
            mismatches.append(f"{f['name']}: report says {f['formatted']}, "
                              f"{f['file']}:{f['column']} re-derives {got}")
        if f["formatted"] not in text:
            mismatches.append(f"{f['name']}: {f['formatted']} is not in the document text")
    return mismatches, sorted(set(notes))


# --------------------------------------------------------------------------
# selftest
# --------------------------------------------------------------------------

def _selftest() -> int:
    checks: list[tuple[str, bool]] = []

    def check(name: str, ok) -> None:
        checks.append((name, bool(ok)))

    def refused(fn, *a, **k) -> bool:
        try:
            fn(*a, **k)
        except rd.ReportError:
            return True
        return False

    import charts as ch

    tmp = tempfile.mkdtemp(prefix="location_report_")
    sha = rd.write_fixture(tmp)
    os.makedirs(os.path.join(tmp, "templates"), exist_ok=True)
    shutil.copy(os.path.join(ROOT, MANIFEST), os.path.join(tmp, MANIFEST))
    shutil.copy(os.path.join(ROOT, ch.PALETTE_JSON), os.path.join(tmp, ch.PALETTE_JSON))
    m = load_manifest(tmp)
    base = brt.to_bytes(brt.build_location(m))
    with open(os.path.join(tmp, BASE_TEMPLATE), "wb") as fh:
        fh.write(base)

    def docx_of(res):
        with open(res["docx"], "rb") as fh:
            return fh.read()

    res = render("13001", root=tmp, sha=sha)
    data = docx_of(res)
    text = plain_text(blocks(data))
    check("enacted county renders with every fact printed",
          all(f["formatted"] in text for f in json.load(open(res["facts"]))["facts"]))
    check("history lists the city moratorium with its level", "Griffin (Alpha County)" in text
          and "| city |" in text)
    check("source links are hyperlinks", b"w:hyperlink" in zipfile.ZipFile(io.BytesIO(data)).read(
        "word/document.xml"))
    defs = rd.load_definitions(tmp)
    check("every term used is defined once, on first use",
          all(text.lower().count(f"{t['term']} ({t['definition']})".lower()) <= 1
              for t in defs["terms"].values()) and "decile (the county" in text)
    check("score chart is embedded", any(n.startswith("word/media/")
                                         for n in zipfile.ZipFile(io.BytesIO(data)).namelist()))
    mism, _ = verify(res["facts"], tmp)
    check("re-derivation: zero mismatches", mism == [])
    md = open(res["md"], encoding="utf-8").read()
    check("markdown twin under deliverables/ with links", res["md"].endswith(
        os.path.join("deliverables", "location_reports", "13001_alpha.md"))
        and "](https://example.org/a)" in md and md.startswith("# Alpha County, Georgia"))

    res0 = render("13003", root=tmp, sha=sha, charts=False)
    bl0 = blocks(docx_of(res0))
    hist = section_text(bl0, "Enacted restriction history")
    check("no-enacted county: one plain sentence, no number, no rate",
          hist is not None and hist.strip() == "No enacted restriction is on record for "
          "Beta County, Georgia." and not rd.DIGIT_RE.search(hist))
    check("no-enacted summary defines the term in its plain sentence",
          "No enacted restriction (a moratorium" in plain_text(bl0))

    # client templates
    cdir = os.path.join(tmp, CLIENT_DIR, "acme")
    os.makedirs(cdir)
    shutil.copy(os.path.join(tmp, BASE_TEMPLATE), os.path.join(cdir, "location_report.docx"))
    rc = render("13001", root=tmp, sha=sha, client="acme", charts=False)
    check("client template selected and named in the output", rc["docx"].endswith("13001_alpha-acme.docx"))
    check("bad client name refused", refused(render, "13001", root=tmp, sha=sha, client="../x"))
    check("missing client template refused", refused(render, "13001", root=tmp, sha=sha, client="zeta"))

    import docx as pydocx

    def client_with(texts: list[str], name: str) -> None:
        d = pydocx.Document()
        for t in texts:
            d.add_paragraph(t)
        os.makedirs(os.path.join(tmp, CLIENT_DIR, name), exist_ok=True)
        d.save(os.path.join(tmp, CLIENT_DIR, name, "location_report.docx"))

    full = ["{{ report.title }}", "{{ define('decile') }}", "{{ plate }}", "{{ history.enacted }}",
            "{{ score.decile }}", "{{ cases.count }}", "{{ neighbors.count }}", "{{ sources.notice }}"]
    client_with(full[:3], "thin")
    check("client template missing required fields refused",
          refused(render, "13001", root=tmp, sha=sha, client="thin", charts=False))
    client_with(full + ["{{ secret }}"], "extra")
    check("client template with an unknown field refused",
          refused(render, "13001", root=tmp, sha=sha, client="extra", charts=False))
    client_with(full + ["{{ report.blocked_share }}"], "grouplevel")
    check("client template reading a refused column refused",
          refused(render, "13001", root=tmp, sha=sha, client="grouplevel", charts=False))
    client_with(full + ["Score of 42 percent"], "typed")
    check("client template with a typed number refused",
          refused(render, "13001", root=tmp, sha=sha, client="typed", charts=False))

    # the plate slot
    png = rd.write_plate(tmp, sha)
    rp = render("13001", root=tmp, sha=sha, plate=png, charts=False)
    blp = blocks(docx_of(rp))
    figs = [i for i, b in enumerate(blp) if b["kind"] == "p" and b["image"]]
    check("plate sits directly under the title block",
          figs and blp[figs[0] - 1]["style"] == "Subtitle")
    check("plate caption and credits come from the sidecar",
          blp[figs[0] + 1]["text"] == "Alpha County has two tracked cases. Alpha County, Georgia."
          and blp[figs[0] + 2]["style"] == "Credit"
          and blp[figs[0] + 2]["text"].startswith("Elevation: U.S. Geological Survey"))
    ext = re.search(rb'<wp:extent cx="(\d+)" cy="(\d+)"', zipfile.ZipFile(
        io.BytesIO(docx_of(rp))).read("word/document.xml"))
    check("plate fills the text width and keeps its aspect ratio",
          ext and int(ext.group(1)) == text_width_emu(base)
          and abs(int(ext.group(1)) * 2 / 3 - int(ext.group(2))) <= 2)
    with zipfile.ZipFile(io.BytesIO(docx_of(rp))) as z:
        media = [z.read(n) for n in z.namelist() if n.startswith("word/media/")]
    check("plate image bytes are inserted unmodified", open(png, "rb").read() in media)
    rn = render("13001", root=tmp, sha=sha, charts=False)
    styles_with = [b.get("style") for b in blp if not (b["kind"] == "p" and (
        b["image"] or b["style"] == "Credit" or b["text"].startswith("Alpha County has two")))]
    styles_without = [b.get("style") for b in blocks(docx_of(rn))]
    check("without --plate the layout is unchanged", styles_with == styles_without)
    write_plate_preview = rd.write_plate(tmp, sha, preview=True)
    check("preview plate refused on a final report",
          refused(render, "13001", root=tmp, sha=sha, plate=write_plate_preview, charts=False))
    check("preview plate accepted on a draft",
          not refused(render, "13001", root=tmp, sha=sha, plate=write_plate_preview, draft=True,
                      charts=False))
    rd.write_plate(tmp, "f" * 40)
    check("plate from a non-ancestor commit refused",
          refused(render, "13001", root=tmp, sha=sha, plate=png, charts=False))
    os.remove(os.path.splitext(png)[0] + ".json")
    check("plate without its sidecar refused",
          refused(render, "13001", root=tmp, sha=sha, plate=png, charts=False))

    # every rendered variant passes the FR-002 rules (post_checks ran) and the
    # re-derivation catches a changed cell
    scores = os.path.join(tmp, rd.SCORES)
    body = open(scores, encoding="utf-8").read().replace("13001,0.2,0.1717,8,1", "13001,0.2,0.2717,9,1")
    with open(scores, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(body)
    mism, notes = verify(res["facts"], tmp)
    check("re-derivation reports a changed cell", any("score.calibrated" in x for x in mism)
          and any("data growth" in n for n in notes))

    fails = 0
    for name, ok in checks:
        print(("PASS " if ok else "FAIL ") + name)
        fails += 0 if ok else 1
    print(f"{len(checks) - fails}/{len(checks)} checks passed")
    return 1 if fails else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--fips")
    ap.add_argument("--client", default="")
    ap.add_argument("--plate")
    ap.add_argument("--draft", action="store_true")
    ap.add_argument("--no-charts", action="store_true")
    ap.add_argument("--allow-uncommitted", action="store_true",
                    help="render from uncommitted inputs (draft preview only; the SHA says so)")
    ap.add_argument("--verify", metavar="FACTS_JSON")
    args = ap.parse_args()
    if args.selftest:
        return _selftest()
    try:
        if args.verify:
            mism, notes = verify(args.verify)
            for n in notes:
                print(f"note: {n}")
            for x in mism:
                print(f"MISMATCH {x}")
            print(f"re-derivation: {len(mism)} mismatch(es)")
            return 1 if mism else 0
        if not args.fips:
            ap.error("--fips, --verify or --selftest is required")
        res = render(args.fips, client=args.client, plate=args.plate, draft=args.draft,
                     charts=not args.no_charts, allow_uncommitted=args.allow_uncommitted)
    except (rd.ReportError, eg.ExportError) as exc:
        print(f"render_location_report: {exc}", file=sys.stderr)
        return 1
    for k in ("docx", "facts", "md"):
        print(f"wrote {os.path.relpath(res[k], ROOT)}")
    print(f"rendered in {res['seconds']} s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
