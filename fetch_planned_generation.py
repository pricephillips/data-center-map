#!/usr/bin/env python3
"""
fetch_planned_generation.py -- planned power plants from EIA-860M.

Why
---
The data center fights of 2026 are increasingly power-plant fights: a campus
that cannot wait for the grid brings its own gas turbines (xAI in Memphis and
Southaven, the Homer City and Shippingport conversions, the Stargate sites).
The generation shows up in EIA's monthly generator inventory, Form 860M, as a
planned unit, often under the developer's or a special-purpose entity's name,
months before the campus is built. data/proposals.csv records whether a
project brings its own power only as a tri-state flag, when it records it at
all. The 860M planned list is the federal record of the plants themselves.

What it does
------------
Downloads the most recent EIA-860M workbook (trying the current month, then
walking back), reads the Planned sheet, and aggregates generator units to one
row per plant: capacity, technology and fuel mix, balancing authority,
earliest planned online date, and coordinates. Plants whose name or owner
reads as data-center-related are flagged, as a prompt for review; the real
link is made spatially in proposal_enrichment.py, which measures planned
generation within reach of every proposed project.

Schema tolerance: EIA moves the header row and renames columns between
vintages. The header row is found by content (the row that carries both a
plant-name and a capacity column), and each concept is resolved through a
candidate list. Missing required concepts abort without writing.

Output
  data/grid_planned_generation.csv
  data/grid_planned_generation_manifest.json

Usage
  python fetch_planned_generation.py --fetch
  python fetch_planned_generation.py --from-file eia860m.xlsx
  python fetch_planned_generation.py --selftest

Public domain (EIA). Requires openpyxl for --fetch / --from-file only.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import tempfile
import urllib.request
from collections import Counter, defaultdict
from datetime import date

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT_CSV = os.path.join(ROOT, "data", "grid_planned_generation.csv")
OUT_MANIFEST = os.path.join(ROOT, "data", "grid_planned_generation_manifest.json")
BASE = "https://www.eia.gov/electricity/data/eia860m/xls"
MONTHS = ["january", "february", "march", "april", "may", "june", "july",
          "august", "september", "october", "november", "december"]
USER_AGENT = "hawthorn-dc-tracker/1.0 (grid context; contact repo owner)"

CONCEPTS = {
    "entity": ["entity name", "utility name", "owner name"],
    "plant_id": ["plant id", "plant code"],
    "plant": ["plant name"],
    "state": ["plant state", "state"],
    "county": ["county"],
    "ba": ["balancing authority code", "balancing authority"],
    "gen_id": ["generator id"],
    "mw": ["nameplate capacity (mw)", "nameplate capacity", "net summer capacity (mw)"],
    "technology": ["technology"],
    "fuel": ["energy source code", "energy source"],
    "month": ["planned operation month", "planned operating month"],
    "year": ["planned operation year", "planned operating year"],
    "status": ["status"],
    "lat": ["latitude"],
    "lon": ["longitude"],
}
REQUIRED = ("plant_id", "plant", "state", "mw")

# A name or owner that reads as data-center-related. A prompt for review, not
# a classification: the spatial join in proposal_enrichment.py is the link.
DC_PATTERN = re.compile(
    r"\b(data ?cent(er|re)s?|hyperscale|digital|compute|ai campus|stargate|"
    r"xai|meta|google|amazon|aws|microsoft|oracle|openai|crusoe|coreweave|"
    r"vantage|qts|cyrusone|aligned|stack infrastructure|edgecore|tract|"
    r"homer city|shippingport|colossus|behind[- ]the[- ]meter)\b", re.I)

FUEL_GROUP = {"NG": "gas", "SUN": "solar", "WND": "wind", "MWH": "storage",
              "NUC": "nuclear", "WAT": "hydro", "DFO": "oil", "BIT": "coal",
              "SUB": "coal", "LFG": "gas", "OBG": "gas", "GEO": "geothermal"}

FIELDS = ["plant_id", "plant_name", "entity", "state", "county", "ba_code",
          "lat", "lon", "planned_mw", "n_units", "fuel_mix", "primary_fuel",
          "technologies", "earliest_online", "statuses", "dc_named"]


def find_header(rows):
    """Index of the header row: the first row carrying plant name and capacity."""
    for i, r in enumerate(rows[:30]):
        cells = {str(c or "").strip().lower() for c in r}
        if "plant name" in cells and any(c.startswith("nameplate capacity") for c in cells):
            return i
    return None


def resolve(header):
    low = [str(c or "").strip().lower() for c in header]
    m = {}
    for concept, cands in CONCEPTS.items():
        for c in cands:
            if c in low:
                m[concept] = low.index(c)
                break
    return m, [c for c in REQUIRED if c not in m]


def _num(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def aggregate(rows):
    """Planned-sheet rows (header first) -> plant-level dicts."""
    h = find_header(rows)
    if h is None:
        raise SystemExit("fetch_planned_generation: no header row found in the Planned sheet")
    m, missing = resolve(rows[h])
    if missing:
        raise SystemExit(f"fetch_planned_generation: missing {missing}; header was {rows[h]}")
    g = lambda r, k: (r[m[k]] if k in m and m[k] < len(r) else None)
    plants = defaultdict(lambda: {"mw": 0.0, "units": 0, "fuels": Counter(), "tech": Counter(),
                                  "online": [], "status": Counter()})
    for r in rows[h + 1:]:
        pid = g(r, "plant_id")
        if pid in (None, "") or not str(pid).strip().replace(".0", "").isdigit():
            continue                               # footnotes and blank rows
        pid = str(pid).strip().replace(".0", "")
        p = plants[pid]
        p.update(plant_name=str(g(r, "plant") or "").strip(),
                 entity=str(g(r, "entity") or "").strip(),
                 state=str(g(r, "state") or "").strip(),
                 county=str(g(r, "county") or "").strip(),
                 ba_code=str(g(r, "ba") or "").strip(),
                 lat=g(r, "lat"), lon=g(r, "lon"))
        mw = _num(g(r, "mw")) or 0.0
        p["mw"] += mw
        p["units"] += 1
        fuel = str(g(r, "fuel") or "").strip().upper()
        p["fuels"][FUEL_GROUP.get(fuel, fuel.lower() or "unknown")] += mw
        if g(r, "technology"):
            p["tech"][str(g(r, "technology")).strip()] += 1
        y, mo = _num(g(r, "year")), _num(g(r, "month"))
        if y:
            p["online"].append(f"{int(y)}-{int(mo):02d}" if mo else f"{int(y)}")
        if g(r, "status"):
            p["status"][str(g(r, "status")).strip()] += 1
    out = []
    for pid, p in plants.items():
        fuels = [f for f, _ in p["fuels"].most_common()]
        out.append({
            "plant_id": pid, "plant_name": p["plant_name"], "entity": p["entity"],
            "state": p["state"], "county": p["county"], "ba_code": p["ba_code"],
            "lat": p["lat"] if p["lat"] not in (None, "") else "",
            "lon": p["lon"] if p["lon"] not in (None, "") else "",
            "planned_mw": round(p["mw"], 1), "n_units": p["units"],
            "fuel_mix": "; ".join(f"{f}:{round(p['fuels'][f])}" for f in fuels),
            "primary_fuel": fuels[0] if fuels else "",
            "technologies": "; ".join(t for t, _ in p["tech"].most_common(3)),
            "earliest_online": min(p["online"]) if p["online"] else "",
            "statuses": "; ".join(s for s, _ in p["status"].most_common()),
            "dc_named": "yes" if DC_PATTERN.search(f"{p['plant_name']} {p['entity']}") else "no",
        })
    out.sort(key=lambda r: (-r["planned_mw"], r["plant_id"]))
    return out


def is_xlsx(data):
    """An xlsx is a zip archive; anything else is an error page."""
    return data[:4] == b"PK\x03\x04"


def download(workdir):
    today = date.today()
    y, mth = today.year, today.month
    for _ in range(8):
        name = f"{MONTHS[mth - 1]}_generator{y}.xlsx"
        url = f"{BASE}/{name}"
        dest = os.path.join(workdir, name)
        try:
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(req, timeout=180) as r:
                data = r.read()
            # EIA answers a month it has not published yet with an HTML page
            # and HTTP 200, not a 404. The first live run (2026-09-29) saved
            # that page as the workbook and crashed in openpyxl instead of
            # falling back a month, so only a real xlsx ends the search.
            if is_xlsx(data):
                with open(dest, "wb") as fh:
                    fh.write(data)
                return dest, name
        except Exception:
            pass
        mth -= 1
        if mth == 0:
            y, mth = y - 1, 12
    raise SystemExit("fetch_planned_generation: no EIA-860M workbook found in the last 8 months")


def read_planned(path):
    import openpyxl                                   # CI-only dependency
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.strip().lower().startswith("planned")), None)
    if not sheet:
        raise SystemExit(f"fetch_planned_generation: no Planned sheet in {wb.sheetnames}")
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def write(plants, vintage):
    with open(OUT_CSV, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(plants)
    with open(OUT_MANIFEST, "w", encoding="utf-8", newline="\n") as fh:
        json.dump({"generated": date.today().isoformat(), "source": "EIA Form 860M",
                   "license": "public domain", "vintage_file": vintage,
                   "plants": len(plants),
                   "planned_mw_total": round(sum(p["planned_mw"] for p in plants)),
                   "dc_named_plants": sum(1 for p in plants if p["dc_named"] == "yes")},
                  fh, indent=1)
        fh.write("\n")


def selftest():
    checks = []

    def check(label, ok):
        checks.append(ok)
        print(f"{'PASS' if ok else 'FAIL'}  {label}")

    rows = [["U.S. Energy Information Administration"], [None],
            ["Entity ID", "Entity Name", "Plant ID", "Plant Name", "Plant State", "County",
             "Balancing Authority Code", "Generator ID", "Nameplate Capacity (MW)",
             "Technology", "Energy Source Code", "Planned Operation Month",
             "Planned Operation Year", "Status", "Latitude", "Longitude"],
            [1, "xAI Power LLC", 900, "Southaven Turbines", "MS", "DeSoto", "MISO", "1", 35.0,
             "Natural Gas Fired Combustion Turbine", "NG", 6, 2026, "(V) Under construction", 34.97, -90.0],
            [1, "xAI Power LLC", 900, "Southaven Turbines", "MS", "DeSoto", "MISO", "2", 35.0,
             "Natural Gas Fired Combustion Turbine", "NG", 3, 2026, "(V) Under construction", 34.97, -90.0],
            [2, "Sunny Co", 901, "Big Solar", "TX", "Pecos", "ERCO", "PV1", 200,
             "Solar Photovoltaic", "SUN", None, 2027, "(P) Planned", 31.0, -103.0],
            ["NOTE: footnote row"]]
    out = aggregate(rows)
    by = {p["plant_id"]: p for p in out}
    check("the header row is found by content, not position", len(out) == 2)
    check("units aggregate to one plant row", by["900"]["n_units"] == 2 and by["900"]["planned_mw"] == 70.0)
    check("the earliest planned online month is kept", by["900"]["earliest_online"] == "2026-03")
    check("a year without a month stays year-precision", by["901"]["earliest_online"] == "2027")
    check("fuel codes fold to groups", by["900"]["primary_fuel"] == "gas" and by["901"]["primary_fuel"] == "solar")
    check("a data-center-named owner is flagged for review", by["900"]["dc_named"] == "yes")
    check("an ordinary plant is not", by["901"]["dc_named"] == "no")
    check("footnote rows are skipped", "NOTE" not in json.dumps(out))
    try:
        aggregate([["Plant Name", "Nameplate Capacity (MW)"], ["x", 1]])
        check("missing required columns abort", False)
    except SystemExit:
        check("missing required columns abort", True)

    # An unpublished month is an HTML page with HTTP 200: skip it, take the
    # month before. Offline: urlopen is stubbed.
    class _Resp:
        def __init__(self, body):
            self.body = body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return self.body

    this_month = MONTHS[date.today().month - 1]
    real = urllib.request.urlopen
    urllib.request.urlopen = lambda req, timeout=0: _Resp(
        b"<!DOCTYPE html><html>page not found</html>" if this_month in req.full_url
        else b"PK\x03\x04workbook")
    try:
        with tempfile.TemporaryDirectory() as d:
            _, got = download(d)
    finally:
        urllib.request.urlopen = real
    check("an HTML page for an unpublished month is skipped", this_month not in got)
    check("and the previous month's workbook is taken", got.endswith(".xlsx"))
    n = sum(checks)
    print(f"\n{n}/{len(checks)} checks passed")
    return 0 if n == len(checks) else 1


def main():
    ap = argparse.ArgumentParser(description="EIA-860M planned generation")
    ap.add_argument("--fetch", action="store_true")
    ap.add_argument("--from-file")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        raise SystemExit(selftest())
    if a.fetch:
        with tempfile.TemporaryDirectory() as td:
            path, name = download(td)
            plants = aggregate(read_planned(path))
    elif a.from_file:
        name = os.path.basename(a.from_file)
        plants = aggregate(read_planned(a.from_file))
    else:
        ap.print_help()
        return
    write(plants, name)
    print(f"fetch_planned_generation: {len(plants)} planned plants from {name}, "
          f"{sum(1 for p in plants if p['dc_named'] == 'yes')} data-center-named -> {OUT_CSV}")


if __name__ == "__main__":
    main()
